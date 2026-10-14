"""Synthetic capacity workbook with the same quirks as a real one.

Real workbooks are never committed; these fixtures reproduce
its structure: models as columns, category header rows, orange "unconfirmed"
cells, combined and component columns, a merged cell, duplicate row names and
a Summary sheet marking NPI platforms.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from tsf_sizer import db

ORANGE = PatternFill(fill_type="solid", fgColor="FFFF9900")
WHITE = PatternFill(fill_type="solid", fgColor="FFFFFFFF")

HEADERS = [
    "PA-7500",
    "PA-7500 MPC",
    "PA-5550",
    "PA-3430",
    "PA-450R and PA-450R-5G",
    "PA-540",
    "PA-455R-5G",
]

# (category or None, attribute, values per HEADERS column)
ROWS = [
    ("Performance", None, None),
    (
        None,
        "App-ID firewall throughput 64k (appmix)",
        ["1.5 Tbps", "-", 175.0, "29.0 Gbps", "3.2 Gbps", "1.8 Gbps", "3.2 Gbps"],
    ),
    (
        None,
        "Threat prevention throughput 64k (appmix)",
        ["1.44 Tbps", "-", "120 Gbps", "15.0 Gbps", "1.4 Gbps", "1.2 Gbps", "1.8 Gbps"],
    ),
    (
        None,
        "IPSec VPN throughput",
        ["407 Gbps", "-", 100.0, "12.0 Gbps", "2.2 Gbps", "650 Mbps", "650 Mbps"],
    ),
    (None, "Connections per second", [7200000, "-", 1670000, 240000, 48000, 15000, 48000]),
    ("Policy", None, None),
    (None, "Security rulebase", [65000, 65000, 65000, 30000, 2000, 1500, 1500]),
    ("Objects (Addresses & Services)", None, None),
    (None, "Max address entries", [160000, 160000, 160000, 30000, 5000, "4K", 5000]),
    (None, "Max address groups", [80000, 80000, 80000, 15000, 1000, 500, 1000]),
    (None, "Max address groups", [1, 1, 1, 1, 1, 1, 1]),
    ("External Dynamic List (EDL), formerly DBL", None, None),
    (None, "Max number of DNS per system", ["4M", "4M", "4M", "1M", "1M", 50000, "1M"]),
    ("Interfaces", None, None),
    (None, "Traffic - 10/100/1000", ["Based on NCs", "-", "-", "-", 8, 8, 6]),
    (None, "Traffic - SFP+ (10G)", ["-", "-", "-", 10, "-", "-", None]),
    (None, "Traffic - QSFP28 (100G)", ["-", "-", 16, 2, "-", "-", "-"]),
    (None, "Traffic - Mystery (7G)", ["-", "-", "-", "-", "-", "-", "-"]),
    (None, "Max interfaces (ifNet)", [8400, 8400, 8400, 4500, 1024, 1024, 1024]),
    (None, "   Tunnel interfaces", [4069, None, None, 4000, None, None, None]),
    (None, "Maximum aggregates with QOS support", [8, 8, "Merged note", None, None, 3, 3]),
    ("Routing", None, None),
    (
        None,
        "Bidirectional Forwarding Detection (BFD) Sessions",
        [1024, 1024, 1024, 512, "No", "No", "No"],
    ),
    (None, "ECMP", [4, 4, 4, "4 (PAN-296613)", "Yes", "Yes", "Yes"]),
    ("L2 Forwarding", None, None),
    (
        None,
        "ARP table size per device",
        [256000, 256000, 132000, "16000 (PAN-255203)", 6000, 3000, 6000],
    ),
    ("High Availability (HA)", None, None),
    (
        None,
        "Track-IP failure detection",
        ["A/P only", "A/P only", "Yes", "Yes", "Yes", "Yes", "Yes"],
    ),
    (None, "Maximum Virtual Addresses (VIP, VMAC)", ["TBD", "TBD", 4096, 128, 32, 32, 32]),
    ("Virtual Systems", None, None),
    (None, "Max virtual systems", [225, 225, 225, 11, "N/A", "N/A", "N/A"]),
    (None, "Max sessions per virtual system", ["Configurable"] * 7),
]

SUMMARY = [
    ("Platform", "Model", None),
    ("PA-400", "PA-410, PA-450R-5G", None),
    ("PA-3400", "PA-3430", None),
    ("PA-500", "PA-540, PA-520", "NPI "),
    ("PA-5500", "PA-5550", "NPI"),
    ("PA-400 NPI", "PA-455R-5G", "NPI"),
    ("Pending coverage", "PA-505, PA-510", "NPI - To do after reg runs"),
    ("Layer", "Capacity Attribute", "QA Manager"),
    ("L4-L7", "Security rule schedules", "Someone"),
]

# Orange ("unconfirmed") cells: (attribute, header)
ORANGE_CELLS = {("Security rulebase", "PA-3430"), ("Max address entries", "PA-540")}


def build_workbook(path, *, overrides: dict | None = None) -> None:
    """Write the synthetic workbook. `overrides` maps (attribute, header) -> new value."""
    overrides = overrides or {}
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    for row in SUMMARY:
        summary.append(list(row))

    ws = wb.create_sheet("12.1.2 L2- L7")
    ws.append([None, "Legend: Orange Cells not tested or not confirmed yet."] + HEADERS)
    for category, attr, vals in ROWS:
        if attr is None:
            ws.append([category])
            continue
        vals = list(vals)
        for (a, h), v in overrides.items():
            if a == attr.strip():
                vals[HEADERS.index(h)] = v
        ws.append([None, attr] + vals)
        r = ws.max_row
        for i, h in enumerate(HEADERS):
            cell = ws.cell(r, 3 + i)
            cell.fill = ORANGE if (attr, h) in ORANGE_CELLS else WHITE
        if attr == "Maximum aggregates with QOS support":
            col = 3 + HEADERS.index("PA-5550")
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)

    wb.create_sheet("ChangeControl").append(["Change Date", "By", "What", "Why"])
    wb.save(path)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "Capacity_Workbook_Test.xlsx"
    build_workbook(path)
    return path


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "app.db")
    yield c
    c.close()


FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def analysis(tmp_path, workbook):
    """(db_path, analysis_id): the synthetic workbook is imported and active, and the sample TSF
    (a PA-3430) is analysed and stored, with its assumptions ledger."""
    import json

    from tsf_sizer import pipeline
    from tsf_sizer.agent import ledger
    from tsf_sizer.portfolio.importer import import_workbook
    from tsf_sizer.sizing.engine import SizingParams

    path = tmp_path / "agent.db"
    c = db.connect(path)
    import_workbook(c, workbook, activate=True)
    params = SizingParams(growth_pct_per_year=20, years=3)
    result = pipeline.analyze(
        c, FIXTURES / "techsupport_sample.txt", params, FIXTURES / "config_sample.xml"
    )
    with c:
        c.execute(
            "INSERT INTO analysis (id, filename, status, params_json, result_json) "
            "VALUES (1, 'sample.txt', 'done', ?, ?)",
            (json.dumps(params.__dict__), json.dumps(result, default=str)),
        )
        ledger.seed(c, 1, params)
    c.close()
    return str(path), 1


@pytest.fixture(autouse=True)
def _no_retry_sleep(monkeypatch):
    from tsf_sizer.agent import gateway as gw

    monkeypatch.setattr(gw, "RETRY_BASE_S", 0.0)


@pytest.fixture
def gateway():
    from fake_gateway import FakeGateway

    gw = FakeGateway()
    yield gw
    gw.close()


@pytest.fixture
def llm_env(gateway, monkeypatch):
    """The environment of a container wired to the (fake) gateway."""
    for key, value in {
        "TSF_SIZER_LLM_MODE": "gateway",
        "TSF_SIZER_GATEWAY_URL": gateway.url,
        "TSF_SIZER_GATEWAY_API_KEY": "secret-key",
        "TSF_SIZER_GATEWAY_MODEL": "fake-model",
        "TSF_SIZER_GATEWAY_MAX_RETRIES": "1",
    }.items():
        monkeypatch.setenv(key, value)
    return gateway
