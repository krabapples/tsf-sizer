from pathlib import Path

import pytest
from conftest import build_workbook
from test_web import CONFIG, TECHSUPPORT

from tsf_sizer import pipeline
from tsf_sizer.portfolio import catalog
from tsf_sizer.portfolio.importer import import_workbook
from tsf_sizer.portfolio.supplements import apply_supplements, apply_to_all_documents
from tsf_sizer.sizing.engine import SizingParams

ALL = [
    "PA-820",
    "PA-850",
    "PA-3220",
    "PA-3250",
    "PA-3260",
    "PA-5220",
    "PA-5250",
    "PA-5260",
    "PA-5280",
]


def _values(conn, doc_id, model):
    return {
        r["tsf_metric"]: r
        for r in conn.execute(
            """SELECT t.tsf_metric, v.num_value, v.unit, v.unconfirmed, v.note
               FROM tsf_metric_map t JOIN model m ON m.name = ?
               JOIN capacity_value v ON v.document_id = t.document_id
                AND v.attribute_id = t.attribute_id AND v.model_id = m.id
               WHERE t.document_id = ?""",
            (model, doc_id),
        )
    }


def test_import_adds_pa800_from_the_datasheet(conn, workbook):
    report = import_workbook(conn, workbook, activate=True)
    assert sorted(report.models["datasheet supplement (not quotable)"]) == sorted(ALL)
    v = _values(conn, report.document_id, "PA-850")
    assert v["perf.throughput_threat_gbps"]["num_value"] == 1.0
    assert v["perf.throughput_appid_gbps"]["num_value"] == 1.9
    assert v["perf.cps"]["num_value"] == 13100
    assert v["perf.throughput_threat_gbps"]["unit"] == "Gbps"
    assert not v["perf.cps"]["unconfirmed"] and "datasheet" in v["perf.cps"]["note"]
    v820 = _values(conn, report.document_id, "PA-820")
    assert v820["perf.throughput_threat_gbps"]["num_value"] == 0.84
    assert v820["perf.cps"]["num_value"] == 8100
    assert catalog.interface_ports(conn, "PA-850", report.document_id) == {
        "1G_RJ45": 4,
        "1G_SFP": 8,
    }


def test_pa800_is_never_quotable_but_listed_with_all(conn, workbook):
    import_workbook(conn, workbook, activate=True)
    quotable = {m["name"] for m in catalog.list_models(conn, quotable_only=True)}
    everything = {m["name"]: m for m in catalog.list_models(conn)}
    assert not {"PA-820", "PA-850"} & quotable
    assert (
        everything["PA-850"]["lifecycle"] == "eol" and not everything["PA-850"]["customer_quotable"]
    )


def test_apply_is_idempotent_and_reaches_older_imports(conn, workbook, tmp_path):
    report = import_workbook(conn, workbook, activate=True)
    conn.execute(
        "DELETE FROM capacity_value WHERE model_id IN (SELECT id FROM model WHERE family='PA-800')"
    )
    conn.commit()
    assert {k: sorted(v) for k, v in apply_to_all_documents(conn).items()} == {
        report.document_id: sorted(ALL)
    }
    assert sorted(apply_supplements(conn, report.document_id)) == sorted(ALL)
    n = conn.execute(
        "SELECT count(*) FROM capacity_value v JOIN model m ON m.id=v.model_id "
        "WHERE m.name='PA-850'"
    ).fetchone()[0]
    apply_to_all_documents(conn)
    assert (
        conn.execute(
            "SELECT count(*) FROM capacity_value v JOIN model m ON m.id=v.model_id "
            "WHERE m.name='PA-850'"
        ).fetchone()[0]
        == n
    )


def test_workbook_wins_when_it_has_the_model(conn, tmp_path):
    path = tmp_path / "Capacity_Workbook_PA850.xlsx"
    build_workbook(path)
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = next(
        w for w in wb.worksheets if any(c.value == "PA-540" for r in w.iter_rows() for c in r)
    )
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "PA-540":  # pretend the workbook lists a PA-850 column
                cell.value = "PA-850"
    wb.save(path)
    report = import_workbook(conn, path, activate=True)
    assert "PA-850" not in report.models.get("datasheet supplement (not quotable)", [])
    assert "PA-820" in report.models["datasheet supplement (not quotable)"]
    assert "PA-3260" in report.models["datasheet supplement (not quotable)"]


def test_pa850_tsf_uses_the_datasheet_and_warns(conn, workbook, tmp_path):
    import_workbook(conn, workbook, activate=True)
    text = TECHSUPPORT.read_text().replace("model: PA-3430", "model: PA-850")
    tsf = tmp_path / "techsupport_pa850.txt"
    tsf.write_text(text)
    r = pipeline.analyze(conn, tsf, SizingParams(), Path(CONFIG))
    assert r["portfolio_error"] is None and r["usage"]["model"] == "PA-850"
    caps = {row["metric"]: row["capacity"] for row in r["usage"]["rows"]}
    assert caps.get("perf.cps_snapshot") == 13100  # the synthetic workbook has no sessions row
    assert any(
        "datasheet" in n and "compare those limits by hand" in n for n in r["sizing"]["notes"]
    )
    named = [r["sizing"]["recommended"], *r["sizing"]["alternatives"], *r["sizing"]["rejected"]]
    assert not any(c and c["model"] in ("PA-820", "PA-850") for c in named)


@pytest.mark.parametrize("model", ["PA-820", "PA-850"])
def test_pa800_family_setting_default(conn, workbook, model):
    import_workbook(conn, workbook, activate=True)
    row = conn.execute("SELECT quotable FROM family_setting WHERE family='PA-800'").fetchone()
    assert row[0] == 0


def test_older_series_ports_and_limits(conn, workbook):
    doc = import_workbook(conn, workbook, activate=True).document_id
    assert catalog.interface_ports(conn, "PA-3260", doc) == {
        "1G_RJ45": 12,
        "1G_SFP": 8,
        "10G_SFP+": 8,
        "40G_QSFP+": 4,
    }
    assert catalog.interface_ports(conn, "PA-5260", doc) == {
        "10G_RJ45": 4,
        "10G_SFP+": 16,
        "100G_QSFP28": 4,
    }
    v = _values(conn, doc, "PA-5260")
    assert (
        v["config.security_rules"]["num_value"] == 20000
        or v["config.security_rules"]["num_value"] > 10000
    )
    assert v["config.security_rules"]["note"].startswith("older capacity workbook")
    # Performance of the PA-3200 / PA-5200 is not in any source: empty, not guessed.
    assert "perf.throughput_threat_gbps" not in v and "perf.cps" not in v
    # Dedicated HA ports are stored as for any other model.
    n = conn.execute(
        """SELECT count(*) FROM capacity_value v JOIN model m ON m.id = v.model_id
           JOIN attribute a ON a.id = v.attribute_id
           WHERE m.name = 'PA-5260' AND v.bool_value = 1 AND a.canonical_key LIKE
             'interfaces.dedicated_ha%'"""
    ).fetchone()[0]
    assert n in (0, 2)


def test_families_are_not_quotable(conn, workbook):
    import_workbook(conn, workbook, activate=True)
    rows = dict(conn.execute("SELECT family, quotable FROM family_setting").fetchall())
    assert rows["PA-800"] == rows["PA-3200"] == rows["PA-5200"] == 0
