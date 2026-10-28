"""Read what the customer sent: PDF files with text and screenshots (for instance of the
Panorama GUI with performance figures), screenshots themselves, and Word, Excel, PowerPoint,
OpenDocument and plain text files.

Text is read from the file. Pictures (PDF pages with pictures, image files, pictures inside
Word and PowerPoint files) are read with OCR (tesseract, installed in the Docker image) so
figures in screenshots are not lost. Everything happens inside the container. OCR can misread
digits: the agent is told to confirm the figures that decide the model.
"""

from __future__ import annotations

import io
import json
import re
import shutil
import sqlite3
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_raw
from PIL import Image

MAX_PAGES = 60
MAX_DOCUMENTS = 10
OCR_DPI = 200
OCR_TIMEOUT_S = 90
OCR_TOTAL_S = 600
MAX_IMAGE_PIXELS = 60_000_000
EXTENSIONS = {
    ".pdf": "pdf", ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx", ".odt": "odt", ".ods": "ods",
    ".txt": "text", ".csv": "text", ".log": "text", ".md": "text",
}  # fmt: skip
LEGACY = {
    ".doc": "an old Word file: open it in Word and save it as .docx (or as a PDF)",
    ".xls": "an old Excel file: save it as .xlsx (or as a PDF)",
    ".ppt": "an old PowerPoint file: save it as .pptx (or as a PDF)",
    ".rtf": "an RTF file: save it as .docx or a PDF",
    ".odp": "an OpenDocument presentation: save it as .pptx or a PDF",
}
KIND_LABELS = {
    "pdf": "PDF", "image": "image", "docx": "Word", "xlsx": "Excel", "pptx": "PowerPoint",
    "odt": "OpenDocument text", "ods": "OpenDocument spreadsheet", "text": "text file",
}  # fmt: skip
PAGE_CHARS = 4000  # a long text is split into pages of about this size
MAX_MEMBER_BYTES = 50_000_000
MAX_ROWS = 3000
MAX_COLUMNS = 60
MAX_PICTURES = 20

Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS


class DocumentError(ValueError):
    """The file cannot be read; the message is shown to the engineer."""


@dataclass
class PageText:
    number: int
    text: str
    method: str  # text | ocr | text+ocr | empty


@dataclass
class Extracted:
    kind: str
    pages: list[PageText] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def chars(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def method(self) -> str:
        methods = {p.method for p in self.pages if p.method != "empty"}
        if not methods:
            return "empty"
        return "+".join(sorted({m for x in methods for m in x.split("+")}))


def ocr_status() -> dict:
    exe = shutil.which("tesseract")
    if not exe:
        return {"available": False, "version": None}
    try:
        out = subprocess.run([exe, "--version"], capture_output=True, text=True, timeout=10)
        version = (out.stdout or out.stderr).splitlines()[0].strip()
    except (OSError, subprocess.SubprocessError, IndexError):
        version = "tesseract"
    return {"available": True, "version": version}


def detect(path: Path, filename: str) -> str:
    """The kind of file, from its content (the extension only has to be plausible)."""
    with path.open("rb") as fh:
        head = fh.read(8)
    ext = Path(filename).suffix.lower()
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG\r\n\x1a\n") or head.startswith(b"\xff\xd8\xff"):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as z:
                names = set(z.namelist())
                mimetype = (
                    z.read("mimetype").decode("ascii", "ignore") if "mimetype" in names else ""
                )
        except (zipfile.BadZipFile, OSError, KeyError):
            raise DocumentError(f"{filename}: the file is damaged") from None
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if "opendocument.text" in mimetype:
            return "odt"
        if "opendocument.spreadsheet" in mimetype:
            return "ods"
        raise DocumentError(f"{filename}: this ZIP file is not a Word, Excel or PowerPoint file")
    if head.startswith(b"\xd0\xcf\x11\xe0"):  # the old binary Office format
        raise DocumentError(
            f"{filename}: "
            + LEGACY.get(ext, "an old Office file: save it as .docx, .xlsx or .pptx")
        )
    if ext in LEGACY:
        raise DocumentError(f"{filename}: " + LEGACY[ext])
    if EXTENSIONS.get(ext) == "text":
        with path.open("rb") as fh:
            sample = fh.read(4096)
        if b"\x00" in sample:
            raise DocumentError(f"{filename}: this is not a text file")
        return "text"
    raise DocumentError(
        f"{filename}: not a supported file. Use PDF, PNG, JPEG, Word (.docx), Excel (.xlsx), "
        "PowerPoint (.pptx), OpenDocument or a text file"
    )


def _ocr(img: Image.Image) -> str:
    exe = shutil.which("tesseract")
    if not exe:
        raise FileNotFoundError("tesseract")
    img = img.convert("RGB")
    if img.width < 1600:  # screenshots are small: bigger letters read better
        factor = 1600 / img.width
        img = img.resize((1600, max(1, int(img.height * factor))), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    out = subprocess.run(
        [exe, "stdin", "stdout", "--psm", "6", "-c", "preserve_interword_spaces=1"],
        input=buf.getvalue(),
        capture_output=True,
        timeout=OCR_TIMEOUT_S,
    )
    return out.stdout.decode("utf-8", errors="replace").strip()


def _clean(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "\n").split("\n")]
    out, blank = [], 0
    for ln in lines:
        blank = blank + 1 if not ln.strip() else 0
        if blank <= 1:
            out.append(ln)
    return "\n".join(out).strip()


def _merge(layer: str, ocr: str) -> str:
    """The text layer, plus the OCR lines that are not already in it."""
    have = {" ".join(ln.split()) for ln in layer.split("\n")}
    extra = [ln for ln in ocr.split("\n") if ln.strip() and " ".join(ln.split()) not in have]
    if not extra:
        return layer
    head = layer + "\n\n" if layer else ""
    return head + "[text read from screenshots on this page]\n" + "\n".join(extra)


def extract(path: Path, filename: str) -> Extracted:
    kind = detect(path, filename)
    if kind in ("docx", "xlsx", "pptx", "odt", "ods", "text"):
        return _extract_office(path, filename, kind)
    ocr_ok = ocr_status()["available"]
    result = Extracted(kind)
    if kind == "image":
        try:
            img = Image.open(path)
            img.load()
        except (OSError, Image.DecompressionBombError) as e:
            raise DocumentError(f"{filename}: the image cannot be read ({e})") from None
        if not ocr_ok:
            result.warnings.append("OCR is not installed: the text in this image was not read.")
            result.pages.append(PageText(1, "", "empty"))
            return result
        try:
            text = _clean(_ocr(img))
        except (OSError, subprocess.SubprocessError) as e:
            raise DocumentError(f"{filename}: OCR failed ({e})") from None
        result.pages.append(PageText(1, text, "ocr" if text else "empty"))
        return result

    try:
        pdf = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError:
        raise DocumentError(f"{filename}: the PDF is damaged or password protected") from None
    total = len(pdf)
    if total > MAX_PAGES:
        result.warnings.append(f"Only the first {MAX_PAGES} of {total} pages were read.")
    spent = 0.0
    import time

    skipped_ocr = False
    for i in range(min(total, MAX_PAGES)):
        page = pdf[i]
        layer = _clean(page.get_textpage().get_text_bounded() or "")
        images = list(page.get_objects(filter=[pdfium_raw.FPDF_PAGEOBJ_IMAGE]))
        needs_ocr = bool(images) or not layer  # a picture, or no text at all
        method = "text" if layer else "empty"
        if needs_ocr and ocr_ok and spent < OCR_TOTAL_S:
            started = time.monotonic()
            try:
                rendered = page.render(scale=OCR_DPI / 72).to_pil()
                ocr_text = _clean(_ocr(rendered))
            except (OSError, subprocess.SubprocessError):
                ocr_text = ""
                result.warnings.append(f"Page {i + 1}: OCR failed.")
            spent += time.monotonic() - started
            merged = _merge(layer, ocr_text)
            if ocr_text:
                method = "text+ocr" if layer else "ocr"
            layer = merged
        elif needs_ocr and images:
            skipped_ocr = True
        result.pages.append(PageText(i + 1, layer, method if layer else "empty"))
    if skipped_ocr:
        result.warnings.append(
            "OCR is not installed or ran out of time: text inside screenshots was not read."
        )
    if all(p.method == "empty" for p in result.pages):
        result.warnings.append("No text found in this file.")
    return result


# --------------------------------------------------------------------------- Office and text


def _member(z: zipfile.ZipFile, name: str) -> bytes:
    info = z.getinfo(name)
    if info.file_size > MAX_MEMBER_BYTES:
        raise DocumentError(f"{name} inside the file is too large")
    return z.read(name)


def _xml(data: bytes, filename: str) -> ET.Element:
    if b"<!ENTITY" in data or b"<!DOCTYPE" in data[:2048] and b"[" in data[:4096]:
        raise DocumentError(f"{filename}: the file contains XML that is not accepted")
    try:
        return ET.fromstring(data)
    except ET.ParseError:
        raise DocumentError(f"{filename}: the file is damaged") from None


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _paragraph_text(p: ET.Element) -> str:
    out = []
    for el in p.iter():
        name = _local(el.tag)
        if name in ("t",) and el.text:
            out.append(el.text)
        elif name == "tab":
            out.append("\t")
        elif name in ("br", "cr"):
            out.append("\n")
    return "".join(out)


def _docx_text(root: ET.Element) -> str:
    """Body text in order; table rows as 'cell | cell'."""
    lines: list[str] = []

    def walk(node: ET.Element) -> None:
        for child in node:
            name = _local(child.tag)
            if name == "p":
                lines.append(_paragraph_text(child))
            elif name == "tbl":
                for row in child.iter():
                    if _local(row.tag) == "tr":
                        cells = []
                        for cell in row:
                            if _local(cell.tag) == "tc":
                                cells.append(
                                    " ".join(
                                        _paragraph_text(p)
                                        for p in cell.iter()
                                        if _local(p.tag) == "p"
                                    ).strip()
                                )
                        lines.append(" | ".join(cells))
            elif name in ("body", "sdt", "sdtContent"):
                walk(child)

    walk(root)
    return "\n".join(lines)


def _pages_from_text(text: str) -> list[str]:
    """Split a long text into pages at line breaks."""
    text = _clean(text)
    if len(text) <= PAGE_CHARS:
        return [text] if text else []
    pages, cur = [], ""
    for line in text.split("\n"):
        while len(line) > PAGE_CHARS:  # one very long line (a CSV without breaks)
            if cur:
                pages.append(cur)
                cur = ""
            pages.append(line[:PAGE_CHARS])
            line = line[PAGE_CHARS:]
        if len(cur) + len(line) + 1 > PAGE_CHARS and cur:
            pages.append(cur)
            cur = ""
        cur += ("\n" if cur else "") + line
    if cur:
        pages.append(cur)
    return pages


def _ocr_picture(data: bytes, name: str, result: Extracted, ocr_ok: bool) -> str:
    """OCR of a picture inside an Office file. Returns the text ('' if none or unavailable)."""
    if not ocr_ok:
        return ""
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
        return _clean(_ocr(img))
    except (OSError, Image.DecompressionBombError, subprocess.SubprocessError):
        result.warnings.append(f"A picture ({name}) could not be read with OCR.")
        return ""


def _extract_office(path: Path, filename: str, kind: str) -> Extracted:
    result = Extracted(kind)
    ocr_ok = ocr_status()["available"]
    texts: list[tuple[str, str]] = []  # (text, method)
    pictures = 0
    if kind == "text":
        raw = path.read_bytes()[:MAX_MEMBER_BYTES]
        text = raw.decode("utf-8-sig", errors="replace")
        texts += [(t, "text") for t in _pages_from_text(text)]
    elif kind == "xlsx":
        import openpyxl

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - openpyxl raises many things for bad files
            raise DocumentError(f"{filename}: the Excel file is damaged") from None
        for ws in wb.worksheets:
            rows = []
            for n, row in enumerate(ws.iter_rows(values_only=True), 1):
                if n > MAX_ROWS:
                    result.warnings.append(
                        f"Sheet '{ws.title}': only the first {MAX_ROWS} rows were read."
                    )
                    break
                cells = ["" if v is None else str(v) for v in row[:MAX_COLUMNS]]
                if any(c.strip() for c in cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                chunks = _pages_from_text(f"Sheet: {ws.title}\n" + "\n".join(rows))
                texts += [(c, "text") for c in chunks]
        wb.close()
        result.warnings.append("Charts and pictures inside the workbook are not read.")
    else:
        try:
            z = zipfile.ZipFile(path)
        except zipfile.BadZipFile:
            raise DocumentError(f"{filename}: the file is damaged") from None
        with z:
            names = z.namelist()
            if kind == "docx":
                root = _xml(_member(z, "word/document.xml"), filename)
                texts += [(t, "text") for t in _pages_from_text(_docx_text(root))]
                media = [n for n in names if n.startswith("word/media/")]
            elif kind == "pptx":
                slides = sorted(
                    (n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                    key=lambda n: int(re.findall(r"\d+", n)[-1]),
                )
                media = []
                for n in slides:
                    root = _xml(_member(z, n), filename)
                    paras = [_paragraph_text(p) for p in root.iter() if _local(p.tag) == "p"]
                    num = re.findall(r"\d+", n)[-1]
                    text = _clean("\n".join(paras))
                    rels = f"ppt/slides/_rels/slide{num}.xml.rels"
                    pics = []
                    if rels in names:
                        for rel in _xml(_member(z, rels), filename):
                            target = rel.get("Target", "")
                            if "media/" in target:
                                pics.append("ppt/media/" + target.rsplit("/", 1)[-1])
                    for pic in pics[:MAX_PICTURES]:
                        if pic in names and pictures < MAX_PICTURES and ocr_ok:
                            pictures += 1
                            ocr_text = _ocr_picture(_member(z, pic), pic, result, ocr_ok)
                            if ocr_text:
                                text = _merge(text, ocr_text)
                    texts.append((f"Slide {num}\n" + text if text else f"Slide {num}", "text"))
                if any(n.startswith("ppt/media/") for n in names) and not ocr_ok:
                    result.warnings.append("OCR is not installed: text in pictures was not read.")
            else:  # odt, ods
                root = _xml(_member(z, "content.xml"), filename)
                lines = []
                if kind == "odt":
                    lines = [_odf_text(p) for p in root.iter() if _local(p.tag) in ("p", "h")]
                else:
                    for row in root.iter():
                        if _local(row.tag) == "table-row":
                            cells = [
                                " ".join(_odf_text(p) for p in c.iter() if _local(p.tag) == "p")
                                for c in row
                                if _local(c.tag) == "table-cell"
                            ]
                            if any(c.strip() for c in cells):
                                lines.append(" | ".join(cells).rstrip(" |"))
                texts += [(t, "text") for t in _pages_from_text("\n".join(lines))]
                media = []
            if kind in ("docx", "odt"):
                for n in media[:MAX_PICTURES]:
                    if not n.lower().endswith((".png", ".jpg", ".jpeg")):
                        continue
                    ocr_text = _ocr_picture(_member(z, n), n, result, ocr_ok)
                    if ocr_text:
                        texts.append(
                            ("[text read from a picture in the document]\n" + ocr_text, "ocr")
                        )
                if media and not ocr_ok:
                    result.warnings.append("OCR is not installed: text in pictures was not read.")
    if len(texts) > MAX_PAGES:
        result.warnings.append(f"Only the first {MAX_PAGES} of {len(texts)} pages were read.")
        texts = texts[:MAX_PAGES]
    for i, (text, method) in enumerate(texts, 1):
        result.pages.append(PageText(i, text, method if text else "empty"))
    if not result.pages:
        result.pages.append(PageText(1, "", "empty"))
        result.warnings.append("No text found in this file.")
    return result


def _odf_text(p: ET.Element) -> str:
    return "".join(p.itertext())


# --------------------------------------------------------------------------- storage


def add_pending(conn: sqlite3.Connection, analysis_id: int, filename: str) -> int:
    return conn.execute(
        "INSERT INTO document (analysis_id, filename, status) VALUES (?, ?, 'reading')",
        (analysis_id, filename[:200]),
    ).lastrowid


def store(conn: sqlite3.Connection, doc_id: int, extracted: Extracted) -> None:
    conn.execute(
        "UPDATE document SET status='done', kind=?, pages=?, chars=?, method=?, pages_json=?, "
        "warnings_json=? WHERE id=?",
        (
            extracted.kind,
            len(extracted.pages),
            extracted.chars,
            extracted.method,
            json.dumps(
                [{"n": p.number, "text": p.text, "method": p.method} for p in extracted.pages]
            ),
            json.dumps(extracted.warnings),
            doc_id,
        ),
    )


def fail(conn: sqlite3.Connection, doc_id: int, message: str) -> None:
    conn.execute(
        "UPDATE document SET status='error', warnings_json=? WHERE id=?",
        (json.dumps([message]), doc_id),
    )


def process(db_path: str, items: list[tuple[int, Path, str]]) -> None:
    """Background job: read each (document id, file, display name). Never raises; the files are
    deleted afterwards."""
    from . import db

    conn = db.connect(db_path)
    try:
        for doc_id, path, name in items:
            try:
                extracted = extract(path, name)
                with conn:
                    store(conn, doc_id, extracted)
            except DocumentError as e:
                with conn:
                    fail(conn, doc_id, str(e))
            except Exception as e:  # noqa: BLE001 - one bad file must not stop the others
                with conn:
                    fail(conn, doc_id, f"{name}: could not be read ({type(e).__name__})")
            finally:
                path.unlink(missing_ok=True)
    finally:
        conn.close()


def list_for(conn: sqlite3.Connection, analysis_id: int) -> list[dict]:
    out = []
    for r in conn.execute(
        "SELECT id, filename, kind, status, pages, chars, method, warnings_json, pages_json "
        "FROM document WHERE analysis_id=? ORDER BY id",
        (analysis_id,),
    ):
        d = dict(r)
        d["warnings"] = json.loads(d.pop("warnings_json") or "[]")
        d["pages_data"] = json.loads(d.pop("pages_json") or "[]")
        out.append(d)
    return out
