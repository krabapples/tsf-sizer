"""Read the capacity workbook into a plain structure.

Layout (12.1.x): one capacity sheet with models as columns and attributes as
rows. Column A holds category headers, column B attribute names, model columns
start at C. A `Summary` sheet lists platform families and marks NPI ones.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# Fill colours that mean "not tested / not confirmed yet" (orange in the legend).
UNCONFIRMED_FILLS = {"FFFF9900"}

_MODEL_HEADER = re.compile(r"^\s*(PA|VM)-", re.I)
_RELEASE = re.compile(r"(\d+\.\d+\.\d+)")


class WorkbookFormatError(ValueError):
    pass


@dataclass
class Cell:
    value: object
    unconfirmed: bool


@dataclass
class AttributeRow:
    sheet_row: int
    category: str
    name: str
    indented: bool
    cells: dict[int, Cell]  # column index -> cell


@dataclass
class SummaryEntry:
    platform: str
    models: list[str]
    status: str


@dataclass
class WorkbookData:
    path: Path
    sha256: str
    sheet_name: str
    release: str | None
    model_columns: dict[int, str]  # column index -> header text
    rows: list[AttributeRow]
    summary: list[SummaryEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _text(v: object) -> str:
    return "" if v is None else str(v)


def _is_unconfirmed(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    rgb = fill.fgColor.rgb if fill.fgColor is not None else None
    return isinstance(rgb, str) and rgb.upper() in UNCONFIRMED_FILLS


def _find_capacity_sheet(wb) -> tuple[Worksheet, int]:
    """Return the sheet and header row that has the most model-name headers."""
    best: tuple[Worksheet, int, int] | None = None
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 10) + 1):
            n = sum(
                1
                for c in range(1, ws.max_column + 1)
                if _MODEL_HEADER.match(_text(ws.cell(r, c).value))
            )
            if n >= 3 and (best is None or n > best[2]):
                best = (ws, r, n)
    if best is None:
        raise WorkbookFormatError("No sheet with model columns (PA-xxxx headers) found")
    return best[0], best[1]


def _merged_values(ws: Worksheet) -> dict[tuple[int, int], object]:
    """Spread the top-left value of every merged range over the whole range."""
    out: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = v
    return out


def _read_summary(wb) -> list[SummaryEntry]:
    if "Summary" not in wb.sheetnames:
        return []
    ws = wb["Summary"]
    entries: list[SummaryEntry] = []
    for r in range(2, ws.max_row + 1):
        platform = _text(ws.cell(r, 1).value).strip()
        if platform.lower() == "layer":  # second table (QA status per attribute) starts here
            break
        models_txt = _text(ws.cell(r, 2).value)
        if not platform and not models_txt:
            continue
        models = [
            re.sub(r"\(.*?\)", "", m).strip()
            for m in models_txt.split(",")
            if re.sub(r"\(.*?\)", "", m).strip()
        ]
        entries.append(SummaryEntry(platform, models, _text(ws.cell(r, 3).value).strip()))
    return entries


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_workbook(path: str | Path) -> WorkbookData:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, header_row = _find_capacity_sheet(wb)
    warnings: list[str] = []

    model_columns: dict[int, str] = {}
    for c in range(1, ws.max_column + 1):
        h = _text(ws.cell(header_row, c).value).strip()
        if _MODEL_HEADER.match(h):
            model_columns[c] = h
    first_model_col = min(model_columns)
    name_col = first_model_col - 1
    category_col = name_col - 1 if name_col > 1 else None

    merged = _merged_values(ws)
    rows: list[AttributeRow] = []
    category = ""
    for r in range(header_row + 1, ws.max_row + 1):
        cat_txt = _text(ws.cell(r, category_col).value).strip() if category_col else ""
        raw_name = _text(ws.cell(r, name_col).value)
        name = raw_name.strip()
        if cat_txt and not name:
            category = cat_txt
            continue
        if not name:
            continue
        if not category:
            warnings.append(f"Row {r}: attribute '{name}' appears before any category header")
        cells = {}
        for c in model_columns:
            v = merged.get((r, c), ws.cell(r, c).value)
            cells[c] = Cell(v, _is_unconfirmed(ws.cell(r, c)))
        indented = raw_name[:1].isspace()
        rows.append(AttributeRow(r, category or "Uncategorized", name, indented, cells))

    m = _RELEASE.search(ws.title) or _RELEASE.search(path.name.replace("_", "."))
    return WorkbookData(
        path=path,
        sha256=file_sha256(path),
        sheet_name=ws.title,
        release=m.group(1) if m else None,
        model_columns=model_columns,
        rows=rows,
        summary=_read_summary(wb),
        warnings=warnings,
    )
